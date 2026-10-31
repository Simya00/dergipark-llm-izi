#!/usr/bin/env python3
"""
Iki Dilli Alanlari Yerel LLM ile Bolme (pilot)
==============================================
Ayni alanda Turkce oz + Ingilizce abstract yazilmis kayitlari (oai_ayristir.py
"iki_dilli_bolundu") yerel bir modele (Ollama: gemma4, llama3.1) bolduruyor.

GUVENCE: Model metni yeniden yazar; bu yuzden modelin verdigi iki parca art
arda konup (bosluklar yok sayilarak) orijinalle HARF HARF karsilastirilir.
Tek harf farkinda sonuc REDDEDILIR. Kabul edilirse parcalar modelin
yazdigindan degil, ORIJINAL metinden kesilir (orijinal karakterler korunur).

Setler (2026-10-31 karari):
  ayar: daha once elle isaretlenen dil_dogrulama.xlsx'teki 50 iki dilli ornek
        (komut bunlarda denenip duzeltilebilir)
  test: daha once gorulmemis, rastgele 100 iki dilli alan (sabit tohum);
        komut bu sette DEGISTIRILMEZ. Kabul esigi: 100'de >= 98 dogru (D/Y).

Komutlar:
  ornekle   ayar ve test setlerini llm_bol_ayar.jsonl / llm_bol_test.jsonl'e yazar
  calistir  --set ayar|test --model gemma4:latest  -> llm_bol_<set>_<model>.jsonl
            (kesilirse ayni komut kaldigi yerden devam eder)
  excel     --set ... --model ...  -> llm_bol_<set>_<model>.xlsx (D/Y listesi)
  hesapla   --set ... --model ...  -> kapsama, dogruluk, %95 Wilson araligi

Gerekli: Ollama calisiyor olmali (ollama serve), openpyxl (excel/hesapla icin).
"""

import argparse
import json
import random
import re
import time
import unicodedata
import urllib.request
from pathlib import Path

from dil_dogrulama import openpyxl_gerekli, wilson

KLASOR = Path(__file__).resolve().parent
OLLAMA = "http://localhost:11434/api/generate"
TOHUM = 20261031
SICAKLIK = 0      # temperature 0: her seferinde ayni cevap
NUM_CTX = 8192    # modelin tek seferde gordugu en uzun metin (token)

KOMUT = """You are given one field from an academic journal record. It contains a
Turkish abstract and an English abstract written one after the other.

Split it into the Turkish part and the English part.

Rules:
1. Copy the text EXACTLY. Do not correct, translate, reformat, add or
   remove any character, including punctuation and spelling errors.
2. Headings such as "Özet", "Öz", "Abstract" stay in the part they
   belong to.
3. A keywords section goes with its language: "Anahtar Kelimeler: ..."
   with Turkish, "Keywords: ..." with English.

Return only JSON:
{"tr": "<Turkish part>", "en": "<English part>"}

Text:
"""

SEMA = {"type": "object",
        "properties": {"tr": {"type": "string"}, "en": {"type": "string"}},
        "required": ["tr", "en"]}


# ---------------------------------------------------------------- ornekleme
def iki_dilli_alanlar(girdi):
    """makaleler.jsonl'deki her iki dilli alan: {url, etiket, orijinal, kural_tr, kural_en}."""
    alanlar = []
    with open(girdi, encoding="utf-8") as g:
        for satir in g:
            m = json.loads(satir)
            bolunen = {}
            for oz in m["ozler"]:
                if oz["dil_kaynagi"] == "iki_dilli_bolundu":
                    if "orijinal_metin" not in oz:
                        raise SystemExit("makaleler.jsonl'de 'orijinal_metin' yok: oai_ayristir.py'yi "
                                         "yeniden calistirin (2026-10-31 degisikligi).")
                    a = bolunen.setdefault((oz.get("etiket"), oz["orijinal_metin"]), {})
                    a[oz["dil"]] = oz["metin"]
            for (etiket, orijinal), p in bolunen.items():
                alanlar.append({"url": m.get("url"), "etiket": etiket, "orijinal": orijinal,
                                "kural_tr": p.get("tr", ""), "kural_en": p.get("en", "")})
    return alanlar


def ayar_urlleri():
    """dil_dogrulama.xlsx'teki iki dilli orneklerin makale adresleri (link hucresinin hedefi)."""
    openpyxl_gerekli()
    from openpyxl import load_workbook
    ws = load_workbook(KLASOR / "dil_dogrulama.xlsx")["dogrulama"]
    return {r[3].hyperlink.target for r in ws.iter_rows(min_row=2)
            if r[1].value == "İki dilli alan ikiye bölündü" and r[3].hyperlink}


def yaz(yol, satirlar):
    with open(yol, "w", encoding="utf-8") as f:
        for s in satirlar:
            f.write(json.dumps(s, ensure_ascii=False) + "\n")


def oku(yol):
    if not yol.exists():
        return []
    with open(yol, encoding="utf-8") as f:
        return [json.loads(s) for s in f if s.strip()]


def komut_ornekle(arg):
    for s in ("ayar", "test"):
        if (KLASOR / f"llm_bol_{s}.jsonl").exists() and not arg.uzerine_yaz:
            raise SystemExit(f"llm_bol_{s}.jsonl zaten var. Yeniden olusturmak icin --uzerine-yaz.")
    alanlar = iki_dilli_alanlar(arg.girdi)
    ayar_url = ayar_urlleri()
    ayar = [a for a in alanlar if a["url"] in ayar_url]
    kalan = [a for a in alanlar if a["url"] not in ayar_url]  # ayar makaleleri testte OLAMAZ
    test = random.Random(TOHUM).sample(kalan, arg.adet)
    for ad, liste in (("ayar", ayar), ("test", test)):
        yaz(KLASOR / f"llm_bol_{ad}.jsonl", [{"no": i, **a} for i, a in enumerate(liste, 1)])
    print(f"iki dilli alan: {len(alanlar)}; ayar: {len(ayar)} alan ({len(ayar_url)} makale adresi); "
          f"test: {len(test)} (tohum {TOHUM})")


# ---------------------------------------------------------------- model + kontrol
def ollama_sor(model, metin):
    istek = {"model": model, "prompt": KOMUT + metin, "stream": False, "format": SEMA,
             "options": {"temperature": SICAKLIK, "seed": TOHUM, "num_ctx": NUM_CTX}}
    req = urllib.request.Request(OLLAMA, data=json.dumps(istek).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=600) as r:
        return json.loads(r.read())["response"]


def sade(metin):
    """Karsilastirma bicimi: Unicode NFC (ayni harfin iki kodlamasi esitlensin) + tum bosluklar silinir."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", metin))


def orijinalden_kes(orijinal, ilk_parca):
    """ilk_parcanin bosluk-disi karakter sayisi kadar ilerleyip orijinali o noktadan ikiye boler."""
    hedef, sayac = len(sade(ilk_parca)), 0
    for i, ch in enumerate(unicodedata.normalize("NFC", orijinal)):
        if sayac == hedef:
            metin = unicodedata.normalize("NFC", orijinal)
            return metin[:i].strip(), metin[i:].strip()
        if not ch.isspace():
            sayac += 1
    return orijinal.strip(), ""


def kontrol(orijinal, cevap):
    """(durum, tr, en). durum: kabul | reddedildi_degisti | reddedildi_json
    Model yalnizca boler; dogru/yanlis kararini kullanici verir (2026-10-31).
    Tek otomatik red: model metni degistirdiyse (orijinalden kesilemez)."""
    try:
        c = json.loads(cevap)
        tr, en = c["tr"], c["en"]
    except (json.JSONDecodeError, KeyError, TypeError):
        return "reddedildi_json", "", ""
    o = sade(orijinal)
    if sade(tr) + sade(en) == o:      # once Turkce yazilmis
        bas, son = orijinalden_kes(orijinal, tr)
        tr, en = bas, son
    elif sade(en) + sade(tr) == o:    # once Ingilizce yazilmis
        bas, son = orijinalden_kes(orijinal, en)
        en, tr = bas, son
    else:
        return "reddedildi_degisti", tr, en
    return "kabul", tr, en


def cikti_yolu(arg, uzanti):
    return KLASOR / f"llm_bol_{arg.set}_{arg.model.replace(':', '_')}.{uzanti}"


def komut_calistir(arg):
    girdiler = oku(KLASOR / f"llm_bol_{arg.set}.jsonl")
    if not girdiler:
        raise SystemExit(f"llm_bol_{arg.set}.jsonl yok: once 'ornekle'.")
    yol = cikti_yolu(arg, "jsonl")
    bitti = {s["no"] for s in oku(yol)}
    with open(yol, "a", encoding="utf-8") as f:
        for g in girdiler:
            if g["no"] in bitti:
                continue
            bas = time.time()
            cevap = ollama_sor(arg.model, g["orijinal"])
            durum, tr, en = kontrol(g["orijinal"], cevap)
            s = {**g, "model": arg.model, "durum": durum, "llm_tr": tr, "llm_en": en,
                 "ham_cevap": cevap, "sure_sn": round(time.time() - bas, 1)}
            f.write(json.dumps(s, ensure_ascii=False) + "\n")
            f.flush()
            print(f"  {g['no']:>3}/{len(girdiler)}  {durum:<22} {s['sure_sn']:>6} sn")


# ---------------------------------------------------------------- excel + hesap
SORU = ("Bölme doğru mu? TR sütununda YALNIZCA Türkçe öz, EN sütununda YALNIZCA İngilizce öz "
        "olmalı; hiçbir cümle yanlış tarafta ya da eksik olmamalı. D = doğru, Y = yanlış")


def komut_excel(arg):
    openpyxl_gerekli()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    yol = cikti_yolu(arg, "xlsx")
    if yol.exists() and not arg.uzerine_yaz:
        raise SystemExit(f"{yol.name} zaten var (cevaplariniz olabilir). --uzerine-yaz ile yeniden olusur.")
    wb = Workbook()
    ws = wb.active
    ws.title = "dogrulama"
    sutunlar = [("no", 5), ("soru", 30), ("adres", 30), ("model sonucu", 16),
                ("TR parça", 70), ("EN parça", 70), ("cevap (D/Y)", 10), ("açıklama", 30)]
    ws.append([a for a, _ in sutunlar])
    for i, (_, gen) in enumerate(sutunlar, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = gen
        ws.cell(1, i).font = Font(bold=True)
    for s in oku(cikti_yolu(arg, "jsonl")):
        kabul = s["durum"] == "kabul"
        url_tr = s["url"].replace("/en/", "/tr/") if s["url"] else ""  # Turkce sayfa (bkz. 2026-10-27 notu)
        ws.append([s["no"], SORU if kabul else "Model reddetti: işaretlemeye gerek yok", url_tr,
                   s["durum"], s["llm_tr"], s["llm_en"], "", ""])
        ws.cell(ws.max_row, 3).hyperlink = url_tr
    for satir in ws.iter_rows():
        for h in satir:
            h.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"D,Y"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{ws.max_row}")
    ws.freeze_panes = "A2"
    wb.save(yol)
    print(f"{ws.max_row - 1} satir -> {yol}")


def komut_hesapla(arg):
    openpyxl_gerekli()
    from openpyxl import load_workbook
    sonuc = {s["no"]: s for s in oku(cikti_yolu(arg, "jsonl"))}
    ws = load_workbook(cikti_yolu(arg, "xlsx")).active
    cevap = {r[0]: str(r[6] or "").strip().upper()[:1] for r in ws.iter_rows(min_row=2, values_only=True)}
    n = len(sonuc)
    kabul = [no for no, s in sonuc.items() if s["durum"] == "kabul"]
    d = sum(cevap.get(no) == "D" for no in kabul)
    y = sum(cevap.get(no) == "Y" for no in kabul)
    bos = len(kabul) - d - y
    print(f"model: {arg.model}  set: {arg.set}  toplam: {n}")
    for durum in ("kabul", "reddedildi_degisti", "reddedildi_json"):
        print(f"  {durum:<22} {sum(s['durum'] == durum for s in sonuc.values()):>4}")
    sure = [s["sure_sn"] for s in sonuc.values()]
    print(f"  ortalama sure: {sum(sure) / len(sure):.1f} sn/alan")
    if bos:
        print(f"  UYARI: {bos} kabul edilen satir isaretlenmemis")
    for ad, dogru, payda in (("kabul edilenlerde dogruluk", d, d + y),
                             ("tumunde (reddedilen = yanlis)", d, n)):
        if payda:
            alt, ust = wilson(dogru, payda)
            print(f"  {ad:<30} {dogru}/{payda} = %{100 * dogru / payda:.1f}  (%95: {100 * alt:.1f}–{100 * ust:.1f})")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("komut", choices=["ornekle", "calistir", "excel", "hesapla"])
    p.add_argument("--set", choices=["ayar", "test"], default="ayar")
    p.add_argument("--model", default="gemma4:latest")
    p.add_argument("--girdi", default=str(KLASOR / "makaleler.jsonl"))
    p.add_argument("--adet", type=int, default=100)
    p.add_argument("--uzerine-yaz", action="store_true")
    arg = p.parse_args()
    {"ornekle": komut_ornekle, "calistir": komut_calistir,
     "excel": komut_excel, "hesapla": komut_hesapla}[arg.komut](arg)


if __name__ == "__main__":
    main()
